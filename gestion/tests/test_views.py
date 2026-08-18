import re
from datetime import date
from decimal import Decimal
from urllib.parse import parse_qs, urlsplit

import pytest
from io import BytesIO
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.utils import timezone
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from asociados.importers import PADRON_IMPORT_SESSION_KEY
from asociados.models import Asociado, CicloLectivo, Curso
from asociados.services import create_asociado
from cuotas.models import Cuota, Donacion, Pago, PagoCuota, PeriodoCuota
from cuotas.services import generar_cuotas_para_periodo, registrar_pago
from gestion.permissions import (
    GESTION_COBRAR_CUOTAS,
    GESTION_CONSULTAR_ASOCIADOS,
    GESTION_EDITAR_ASOCIADOS,
    GESTION_IMPORTAR_ASOCIADOS,
    GESTION_IMPORTAR_CUOTAS_HISTORICAS,
    GESTION_PERMISSIONS,
)
from usuarios.services import ASOCIADO_GROUP
from usuarios.roles import ADMINISTRADOR_APP_GROUP


def crear_usuario_gestion(username="usuario_gestion", permisos=None):
    user_model = get_user_model()
    user = user_model.objects.create_user(username=username, password="secreto123")
    codenames = [permission.split(".", 1)[1] for permission in (permisos or GESTION_PERMISSIONS)]
    user.user_permissions.add(*Permission.objects.filter(content_type__app_label="gestion", codename__in=codenames))
    return user


def crear_planilla_padron(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PADRÓN GENERAL"
    worksheet.append(
        [
            None,
            "Apellido/nombre",
            "Curso/división/Ciclo/turno",
            "CICLO",
            "Categoria (act./adh.)",
            "DNI",
            "Celular",
            "Mail",
            "Dirección",
        ]
    )
    for row in rows:
        worksheet.append(row)
    from io import BytesIO

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "padron.xlsx"
    return buffer


def crear_planilla_cuotas(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "COBRO CUOTAS SOCIALES"
    worksheet.append([None, None, None, None, None, "PAGO DE CUOTAS SOCIALES 2026"])
    worksheet.append([])
    worksheet.append([None, "ACTIVOS", None, "A/T: pago a término del 1 al 10 de cada mes"])
    worksheet.append([None, "ADHERENTES", None, "C/I: pago con interés después del 10"])
    worksheet.append([None, None, None, "AFT: asociado fuera de término (no se asoció en marzo)"])
    worksheet.append([])
    worksheet.append(
        [
            "Socio N°",
            "Apellido/Nombre",
            "curso/división",
            "Mar",
            "forma pago",
            "Abr",
            "forma pago",
            "May",
            "forma pago",
            "Jun",
            "froma pago",
        ]
    )
    for row in rows:
        worksheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "cuotas.xlsx"
    return buffer


@pytest.fixture
def operacion_cuotas_historicas_a_junio(monkeypatch):
    """Mantiene determinista el escenario de la planilla histórica marzo-junio."""
    monkeypatch.setattr("gestion.views.timezone.localdate", lambda: date(2026, 6, 30))


@pytest.mark.django_db
def test_ruta_anterior_del_dashboard_gestion_fue_retirada(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="aso_sin_staff", password="secreto123")
    client.force_login(user)

    response = client.get("/gestion/")

    assert response.status_code == 404


@pytest.mark.django_db
def test_home_gestion_muestra_accesos_basicos(client):
    staff = crear_usuario_gestion("staff1")

    client.force_login(staff)
    response = client.get(reverse("web:home"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Panel de gestión" in content
    assert "Atención al asociado" in content
    assert "Períodos de cuota" in content
    assert "asociados activos" not in content.lower()


@pytest.mark.django_db
def test_home_atencion_muestra_solo_operacion_diaria(client):
    atencion = crear_usuario_gestion(
        "atencion_dashboard",
        permisos=[
            GESTION_CONSULTAR_ASOCIADOS,
            GESTION_EDITAR_ASOCIADOS,
            GESTION_COBRAR_CUOTAS,
        ],
    )

    client.force_login(atencion)
    response = client.get(reverse("web:home"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Atención al asociado" in content
    assert "Importar padrón inicial" not in content
    assert "Importar cuotas históricas" not in content


@pytest.mark.django_db
def test_home_gestion_muestra_importaciones_autorizadas(client):
    admin_operativo = crear_usuario_gestion(
        "admin_importaciones",
        permisos=[
            GESTION_IMPORTAR_ASOCIADOS,
            GESTION_IMPORTAR_CUOTAS_HISTORICAS,
        ],
    )

    client.force_login(admin_operativo)
    response = client.get(reverse("web:home"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Importar padrón inicial" in content
    assert "Importar cuotas históricas" in content


@pytest.mark.django_db
def test_grupo_administrador_app_sin_superusuario_no_habilita_importacion(client):
    user = get_user_model().objects.create_user(username="app_sin_super", password="secreto123")
    user.groups.add(Group.objects.get(name=ADMINISTRADOR_APP_GROUP))

    client.force_login(user)
    response = client.get(reverse("gestion:importar_asociados"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_superusuario_administrador_app_puede_abrir_importacion(client):
    user = get_user_model().objects.create_superuser(
        username="app_super", password="secreto123", email="app@example.com"
    )
    user.groups.add(Group.objects.get(name=ADMINISTRADOR_APP_GROUP))

    client.force_login(user)
    response = client.get(reverse("gestion:importar_asociados"))

    assert response.status_code == 200


@pytest.mark.django_db
def test_cobros_gestion_requiere_permiso(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="no_staff", password="secreto123")
    client.force_login(user)

    response = client.get(reverse("gestion:cobros"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_deudores_gestion_requiere_permiso(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="no_staff_deuda", password="secreto123")
    client.force_login(user)

    response = client.get(reverse("gestion:deudores"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_periodos_cuota_gestion_requiere_permiso(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="no_staff_periodos", password="secreto123")
    client.force_login(user)

    response = client.get(reverse("gestion:periodos_cuota"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_asociados_gestion_requiere_permiso(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="no_staff_asoc", password="secreto123")
    client.force_login(user)

    response = client.get(reverse("gestion:asociados"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_exportar_asociados_requiere_permiso(client):
    user_model = get_user_model()
    user = user_model.objects.create_user(username="no_staff_exporta", password="secreto123")
    client.force_login(user)

    response = client.get(reverse("gestion:exportar_asociados"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_importar_asociados_previsualiza_y_guarda_en_sesion(client):
    staff = crear_usuario_gestion("staff_importa")
    archivo = crear_planilla_padron(
        [
            [1, "Leyes Lena Muriel", "1°2°", "CB", "Activo", "52328996", "2984 111111", "lena@example.com", "Calle 1"],
            [2, "Joaquin Darosa", "1ro C.B", "CB", "Activo", "52536191", "2984 222222", "joaquin@example.com", "Calle 2"],
            [3, "", "", "", "", "", "", "", ""],
        ]
    )

    client.force_login(staff)
    response = client.post(
        reverse("gestion:importar_asociados"),
        {"action": "preview", "archivo": archivo},
        follow=True,
    )

    assert response.status_code == 200
    preview = client.session[PADRON_IMPORT_SESSION_KEY]
    assert preview["summary"]["importar"] == 1
    assert preview["summary"]["revisar"] == 1
    assert preview["summary"]["cursos_a_crear"] == 1
    assert preview["cursos_a_crear"][0]["curso"] == "1ro 2da CB TM"
    content = response.content.decode()
    assert "Cursos que se crearían" in content
    assert "Falta división/comisión del curso" in content
    assert "Cómo leer la previsualización" in content
    assert "Descargar planilla para revisar" in content


@pytest.mark.django_db
def test_importar_asociados_ordena_cursos_por_ciclo_anio_division_y_turno(client):
    staff = crear_usuario_gestion("staff_orden_cursos")
    archivo = crear_planilla_padron(
        [
            [1, "Perez Ana", "4°2°", "CS", "Activo", "42328996", "2984 111111", "ana@example.com", "Calle 1"],
            [2, "Lopez Beto", "1°2°", "CB", "Activo", "42536191", "2984 222222", "beto@example.com", "Calle 2"],
            [3, "Garcia Carla", "4°1°", "CS", "Activo", "43328996", "2984 333333", "carla@example.com", "Calle 3"],
            [4, "Sosa Diego", "1°1°", "CB", "Activo", "43536191", "2984 444444", "diego@example.com", "Calle 4"],
        ]
    )

    client.force_login(staff)
    response = client.post(reverse("gestion:importar_asociados"), {"action": "preview", "archivo": archivo})

    assert response.status_code == 302
    preview = client.session[PADRON_IMPORT_SESSION_KEY]
    assert [curso["curso"] for curso in preview["cursos_a_crear"]] == [
        "1ro 1ra CB TM",
        "1ro 2da CB TM",
        "4to 1ra CS TM",
        "4to 2da CS TM",
    ]


@pytest.mark.django_db
def test_importar_asociados_descarga_planilla_con_filas_a_revisar(client):
    staff = crear_usuario_gestion("staff_descarga_revisar")
    archivo = crear_planilla_padron(
        [
            [1, "Leyes Lena Muriel", "1°2°", "CB", "Activo", "52328996", "2984 111111", "lena@example.com", "Calle 1"],
            [2, "Joaquin Darosa", "1ro C.B", "CB", "Activo", "52536191", "2984 222222", "joaquin@example.com", "Calle 2"],
        ]
    )

    client.force_login(staff)
    client.post(reverse("gestion:importar_asociados"), {"action": "preview", "archivo": archivo})
    response = client.get(reverse("gestion:descargar_asociados_revisar"))

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "padron_asociados_a_revisar.xlsx" in response["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["PADRÓN GENERAL"]
    assert [worksheet.cell(1, column).value for column in range(1, 12)] == [
        "Fila original",
        "Número de asociado",
        "Apellido/nombre",
        "Curso/división/Ciclo/turno",
        "CICLO",
        "Categoria (act./adh.)",
        "DNI",
        "Celular",
        "Mail",
        "Dirección",
        "Motivo de revisión",
    ]
    assert worksheet.max_row == 2
    row_values = [worksheet.cell(2, column).value for column in range(1, 12)]
    assert row_values[:10] == [
        3,
        "2",
        "Joaquin Darosa",
        "1ro C.B",
        "CB",
        "Activo",
        "52536191",
        "2984 222222",
        "joaquin@example.com",
        "Calle 2",
    ]
    assert "Falta división/comisión del curso" in row_values[10]
    assert "curso incompleto o dudoso para asociado" in row_values[10]


@pytest.mark.django_db
def test_importar_asociados_confirma_desde_sesion(client):
    staff = crear_usuario_gestion("staff_confirma")
    archivo = crear_planilla_padron(
        [
            [1, "Leyes Lena Muriel", "1°2°", "CB", "Activo", "52328996", "2984 111111", "lena@example.com", "Calle 1"],
            [2, "Joaquin Darosa", "1ro C.B", "CB", "Activo", "52536191", "2984 222222", "joaquin@example.com", "Calle 2"],
        ]
    )

    client.force_login(staff)
    client.post(reverse("gestion:importar_asociados"), {"action": "preview", "archivo": archivo})
    response = client.post(reverse("gestion:importar_asociados"), {"action": "confirm"}, follow=True)

    assert response.status_code == 200
    assert PADRON_IMPORT_SESSION_KEY not in client.session
    assert Curso.objects.filter(anio="1ro", curso="2da", division=Curso.DIVISION_CB, turno=Curso.TURNO_TM).exists()
    asociado = Asociado.objects.get(dni="52328996")
    assert asociado.apellido == "Leyes"
    assert asociado.nombre == "Lena Muriel"
    assert asociado.numero_asociado == 1
    assert asociado.direccion == "Calle 1"
    assert asociado.fecha_alta == timezone.localdate()
    assert asociado.fecha_inicio_cobro == date(2026, 3, 1)
    assert asociado.usuario is None
    assert not Asociado.objects.filter(dni="52536191").exists()
    assert "1 creados" in response.content.decode()


@pytest.mark.django_db
def test_crear_usuarios_faltantes_asociados_desde_importacion(client):
    staff = crear_usuario_gestion("staff_usuarios_faltantes", permisos=[GESTION_IMPORTAR_ASOCIADOS])
    asociado = Asociado.objects.create(
        nombre="Lena",
        apellido="Leyes",
        dni="52328996",
        tipo=Asociado.TIPO_ASOCIADO,
        fecha_alta="2026-03-01",
        fecha_inicio_cobro="2026-03-01",
    )

    client.force_login(staff)
    response = client.post(reverse("gestion:crear_usuarios_asociados_faltantes"), follow=True)

    asociado.refresh_from_db()
    assert response.status_code == 200
    assert asociado.usuario is not None
    assert asociado.usuario.username == "52328996"
    assert asociado.usuario.check_password("52328996")
    assert Group.objects.get(name=ASOCIADO_GROUP) in asociado.usuario.groups.all()
    assert "Usuarios de asociados creados: 1" in response.content.decode()


@pytest.mark.django_db
def test_crear_usuarios_faltantes_asociados_se_procesa_en_lotes(client, monkeypatch):
    staff = crear_usuario_gestion("staff_usuarios_faltantes_lotes", permisos=[GESTION_IMPORTAR_ASOCIADOS])
    primero = Asociado.objects.create(
        nombre="Lena",
        apellido="Leyes",
        dni="52328996",
        tipo=Asociado.TIPO_ASOCIADO,
        fecha_alta="2026-03-01",
        fecha_inicio_cobro="2026-03-01",
    )
    segundo = Asociado.objects.create(
        nombre="Joaquin",
        apellido="Darosa",
        dni="52536191",
        tipo=Asociado.TIPO_ASOCIADO,
        fecha_alta="2026-03-01",
        fecha_inicio_cobro="2026-03-01",
    )

    monkeypatch.setattr("gestion.views.GestionCrearUsuariosAsociadosFaltantesView.usuarios_batch_size", 1)

    client.force_login(staff)
    response_1 = client.post(reverse("gestion:crear_usuarios_asociados_faltantes"), follow=True)

    primero.refresh_from_db()
    segundo.refresh_from_db()
    assert response_1.status_code == 200
    assert primero.usuario is not None
    assert segundo.usuario is None
    assert "Continuar creando usuarios faltantes" in response_1.content.decode()
    state = client.session["usuarios_asociados_faltantes_state"]
    assert state["cursor"] == primero.id
    assert state["creados"] == 1

    response_2 = client.post(reverse("gestion:crear_usuarios_asociados_faltantes"), follow=True)

    primero.refresh_from_db()
    segundo.refresh_from_db()
    assert response_2.status_code == 200
    assert primero.usuario is not None
    assert segundo.usuario is not None
    assert "Continuar creando usuarios faltantes" not in response_2.content.decode()
    assert "Usuarios de asociados creados: 2" in response_2.content.decode()
    assert "usuarios_asociados_faltantes_state" not in client.session


@pytest.mark.django_db
def test_importar_cuotas_historicas_previsualiza_desde_planilla(client, operacion_cuotas_historicas_a_junio):
    staff = crear_usuario_gestion("staff_cuotas_preview")
    asociado = create_asociado(
        nombre="Lena", apellido="Leyes", dni="52328996", tipo="asociado", fecha_alta="2026-03-01"
    )
    archivo = crear_planilla_cuotas(
        [
            [
                asociado.numero_asociado,
                "Leyes Lena",
                "1°2°",
                True,
                "efectivo",
                False,
                None,
                True,
                "MP",
                False,
                None,
            ],
            [999, "Persona Inexistente", "1°1°", True, "efectivo", False, None, False, None, False, None],
            [1000, None, "1°1°", True, "efectivo", False, None, False, None, False, None],
        ]
    )

    client.force_login(staff)
    response = client.post(
        reverse("gestion:importar_cuotas_historicas"),
        {"action": "preview", "archivo": archivo},
        follow=True,
    )

    assert response.status_code == 200
    preview = client.session["cuotas_historicas_preview"]
    assert preview["summary"]["cuotas_a_crear"] == 4
    assert preview["summary"]["pagos_a_crear"] == 2
    assert preview["summary"]["cuotas_impagas"] == 2
    assert preview["summary"]["revisar"] == 4
    assert preview["summary"]["omitidas"] == 1
    content = response.content.decode()
    assert "Cuotas importables" in content
    assert "No se encontró asociado por nombre" in content
    assert "crea cuotas, pagos y aplicaciones a cuota" in content


@pytest.mark.django_db
def test_importar_cuotas_historicas_descarga_planilla_con_cuotas_a_revisar(
    client, operacion_cuotas_historicas_a_junio
):
    staff = crear_usuario_gestion("staff_cuotas_revisar")
    archivo = crear_planilla_cuotas(
        [
            [999, "Persona Inexistente", "1°1°", True, "efectivo", False, None, False, "MP", False, None],
        ]
    )

    client.force_login(staff)
    client.post(reverse("gestion:importar_cuotas_historicas"), {"action": "preview", "archivo": archivo})
    response = client.get(reverse("gestion:descargar_cuotas_historicas_revisar"))

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "cuotas_historicas_a_revisar.xlsx" in response["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["COBRO CUOTAS SOCIALES"]
    assert [worksheet.cell(1, column).value for column in range(1, 17)] == [
        "Fila original",
        "Número de asociado",
        "Apellido/Nombre",
        "curso/división",
        "Mar pago",
        "Mar forma",
        "Mar motivo",
        "Abr pago",
        "Abr forma",
        "Abr motivo",
        "May pago",
        "May forma",
        "May motivo",
        "Jun pago",
        "Jun forma",
        "Jun motivo",
    ]
    assert worksheet.max_row == 2
    row_values = [worksheet.cell(2, column).value for column in range(1, 17)]
    assert row_values[:6] == [8, 999, "Persona Inexistente", "1°1°", "Sí", "efectivo"]
    assert "No se encontró asociado por nombre" in row_values[6]
    assert row_values[10:12] == ["Sí", "MP"]
    assert "No se encontró asociado por nombre" in row_values[12]


@pytest.mark.django_db
def test_importar_cuotas_historicas_confirma_cuotas_y_pagos(client, operacion_cuotas_historicas_a_junio):
    staff = crear_usuario_gestion("staff_cuotas_confirma")
    asociado = create_asociado(
        nombre="Lena", apellido="Leyes", dni="52328996", tipo="asociado", fecha_alta="2026-03-01"
    )
    archivo = crear_planilla_cuotas(
        [
            [
                asociado.numero_asociado,
                "Leyes Lena",
                "1°2°",
                True,
                "efectivo",
                False,
                None,
                True,
                "MP",
                False,
                None,
            ],
        ]
    )

    client.force_login(staff)
    client.post(reverse("gestion:importar_cuotas_historicas"), {"action": "preview", "archivo": archivo})
    response = client.post(reverse("gestion:importar_cuotas_historicas"), {"action": "confirm"}, follow=True)

    assert response.status_code == 200
    assert PeriodoCuota.objects.filter(ciclo_lectivo__anio=2026, mes__in=[3, 4, 5, 6]).count() == 4
    assert Cuota.objects.filter(asociado=asociado).count() == 4
    assert Pago.objects.filter(asociado=asociado).count() == 2
    assert PagoCuota.objects.filter(pago__asociado=asociado).count() == 2
    marzo = Cuota.objects.get(asociado=asociado, periodo__mes=3)
    abril = Cuota.objects.get(asociado=asociado, periodo__mes=4)
    mayo = Cuota.objects.get(asociado=asociado, periodo__mes=5)
    assert marzo.estado == Cuota.ESTADO_PAGADA
    assert marzo.importe == 500
    assert marzo.importe_recargo_mes == 100
    assert abril.estado == Cuota.ESTADO_VENCIDA
    assert mayo.estado == Cuota.ESTADO_PAGADA
    assert mayo.importe == 800
    assert mayo.importe_recargo_mes == 100
    pago_mayo = Pago.objects.get(asociado=asociado, fecha=date(2026, 5, 10))
    assert pago_mayo.importe == 800
    assert PagoCuota.objects.get(pago=pago_mayo).importe == 800
    assert "4 cuotas creadas" in response.content.decode()


@pytest.mark.django_db
def test_cobros_gestion_usa_consulta_de_asociados_y_registra_pago(client):
    staff = crear_usuario_gestion("staff_cobro")
    asociado = create_asociado(
        nombre="Paula", apellido="Gimenez", dni="45555111", tipo="asociado", fecha_alta="2026-05-10"
    )
    periodo = PeriodoCuota.objects.create(
        mes=timezone.localdate().month,
        ciclo_lectivo=CicloLectivo.objects.get_or_create(anio=timezone.localdate().year)[0],
        importe="3000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
    )
    generar_cuotas_para_periodo(periodo)

    client.force_login(staff)

    response_busqueda = client.get(reverse("gestion:asociados"), {"q": "45555111", "estado": "activo", "usuario": "con"})
    assert response_busqueda.status_code == 200
    content = response_busqueda.content.decode()
    assert "Gimenez" in content
    assert "Estado" in content
    assert "Usuario vinculado" in content
    assert "uni2-row-link" in content
    assert "Ver detalle" not in content
    assert ">Cobrar</a>" not in content
    cuota = asociado.cuotas.get(periodo=periodo)

    response_cobro = client.post(
        reverse("gestion:cobros"),
        {
            "asociado_id": asociado.id,
            "cuotas_ids": [str(cuota.id)],
            "fecha": timezone.localdate().isoformat(),
            "importe": "3000.00",
            "metodo": Pago.METODO_EFECTIVO,
            "observaciones": "Pago en mostrador",
        },
        follow=True,
    )

    assert response_cobro.status_code == 200
    assert Pago.objects.filter(asociado=asociado, importe="3000.00").exists()
    cuota.refresh_from_db()
    assert cuota.estado == cuota.ESTADO_PAGADA
    assert "registrado para Gimenez, Paula" in response_cobro.content.decode()
    assert urlsplit(response_cobro.redirect_chain[-1][0]).path == reverse(
        "gestion:asociado_detalle", args=[asociado.id]
    )


@pytest.mark.django_db
def test_cobros_con_asociado_preseleccionado_no_muestra_busqueda_sin_resultados(client):
    staff = crear_usuario_gestion("staff_cobro_preseleccionado")
    asociado = create_asociado(
        nombre="Paula", apellido="Gimenez", dni="45555112", tipo="asociado", fecha_alta="2026-05-10"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:cobros"), {"asociado": asociado.id})

    assert response.status_code == 200
    content = response.content.decode()
    assert "Gimenez, Paula" in content
    assert "No se encontraron asociados para la búsqueda ingresada" not in content


@pytest.mark.django_db
def test_cobros_sin_asociado_redirige_a_atencion_al_asociado(client):
    staff = crear_usuario_gestion("staff_cobro_sin_asociado")

    client.force_login(staff)
    response = client.get(reverse("gestion:cobros"), follow=True)

    assert response.status_code == 200
    assert response.redirect_chain[-1][0] == reverse("gestion:asociados")
    content = response.content.decode()
    assert "Elegí un asociado desde Atención al asociado" in content
    assert "Atención al asociado" in content


@pytest.mark.django_db
def test_cobros_renderiza_saldo_y_etiqueta_accesible_para_cada_cuota(client):
    staff = crear_usuario_gestion("staff_cobro_saldo_js")
    asociado = create_asociado(
        nombre="Laura", apellido="Mendez", dni="47777112", tipo="asociado", fecha_alta="2026-05-10"
    )
    periodo = PeriodoCuota.objects.create(
        mes=timezone.localdate().month,
        ciclo_lectivo=CicloLectivo.objects.get_or_create(anio=timezone.localdate().year)[0],
        importe="3000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
    )
    generar_cuotas_para_periodo(periodo)
    cuota = asociado.cuotas.get(periodo=periodo)

    client.force_login(staff)
    response = client.get(reverse("gestion:cobros"), {"asociado": asociado.id})

    assert response.status_code == 200
    content = response.content.decode()
    match = re.search(r'data-saldo="([^"]+)"', content)
    assert match
    assert "," not in match.group(1)
    assert Decimal(match.group(1)) == Decimal("3000.00")
    assert f'id="cuota-{cuota.id}"' in content
    assert f'for="cuota-{cuota.id}"' in content
    assert f"Cobrar cuota {cuota.periodo}" in content


@pytest.mark.django_db
def test_cobros_gestion_permite_pago_mayor_y_genera_donacion(client):
    staff = crear_usuario_gestion("staff_donacion")
    asociado = create_asociado(
        nombre="Ivan", apellido="Molina", dni="46666111", tipo="asociado", fecha_alta="2026-05-10"
    )
    periodo = PeriodoCuota.objects.create(
        mes=timezone.localdate().month,
        ciclo_lectivo=CicloLectivo.objects.get_or_create(anio=timezone.localdate().year)[0],
        importe="3000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
    )
    generar_cuotas_para_periodo(periodo)
    cuota = asociado.cuotas.get(periodo=periodo)

    client.force_login(staff)
    response = client.post(
        reverse("gestion:cobros"),
        {
            "asociado_id": asociado.id,
            "cuotas_ids": [str(cuota.id)],
            "fecha": timezone.localdate().isoformat(),
            "importe": "4000.00",
            "metodo": Pago.METODO_EFECTIVO,
            "observaciones": "",
        },
    )

    assert response.status_code == 302
    assert Pago.objects.count() == 1
    pago = Pago.objects.first()
    assert pago.importe == Decimal("4000")
    assert Donacion.objects.filter(pago=pago, importe=Decimal("1000")).exists()


@pytest.mark.django_db
def test_asociado_sin_deuda_puede_registrar_donacion_desde_su_detalle(client):
    staff = crear_usuario_gestion("staff_donacion_sin_deuda")
    asociado = create_asociado(
        nombre="Ana",
        apellido="Paz",
        dni="46666112",
        tipo="asociado",
        fecha_alta="2026-05-10",
    )
    client.force_login(staff)

    detalle = client.get(reverse("gestion:asociado_detalle", args=[asociado.id]))
    formulario = client.get(reverse("gestion:cobros"), {"asociado": asociado.id})
    response = client.post(
        reverse("gestion:cobros"),
        {
            "asociado_id": asociado.id,
            "fecha": timezone.localdate().isoformat(),
            "importe": "2500.00",
            "metodo": Pago.METODO_BILLETERA,
            "observaciones": "Donación voluntaria",
        },
        follow=True,
    )

    assert "Registrar donación" in detalle.content.decode()
    contenido_formulario = formulario.content.decode()
    assert formulario.status_code == 200
    assert "El asociado no tiene cuotas pendientes" in contenido_formulario
    assert "Importe de la donación" in contenido_formulario
    assert "Cuotas a cobrar" not in contenido_formulario
    assert response.status_code == 200
    assert "Donación registrada para Paz, Ana" in response.content.decode()
    pago = Pago.objects.get(asociado=asociado)
    assert not pago.aplicaciones.exists()
    assert Donacion.objects.get(pago=pago).importe == Decimal("2500")


@pytest.mark.django_db
def test_cobros_gestion_permite_cobrar_solo_cuotas_mas_viejas_seleccionadas(client):
    staff = crear_usuario_gestion("staff_cobro_parcial_de_lista")
    asociado = create_asociado(
        nombre="Noelia", apellido="Sosa", dni="48888111", tipo="asociado", fecha_alta="2026-03-10"
    )
    ciclo = CicloLectivo.objects.get_or_create(anio=2026)[0]
    marzo = PeriodoCuota.objects.create(
        mes=3,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-03-10",
    )
    abril = PeriodoCuota.objects.create(
        mes=4,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-04-10",
    )
    generar_cuotas_para_periodo(marzo)
    generar_cuotas_para_periodo(abril)
    cuota_marzo, cuota_abril = asociado.cuotas.order_by("periodo__mes")

    client.force_login(staff)
    response = client.post(
        reverse("gestion:cobros"),
        {
            "asociado_id": asociado.id,
            "cuotas_ids": [str(cuota_marzo.id)],
            "fecha": "2026-04-05",
            "importe": "4000.00",
            "metodo": Pago.METODO_EFECTIVO,
            "observaciones": "",
        },
    )

    assert response.status_code == 302
    cuota_marzo.refresh_from_db()
    cuota_abril.refresh_from_db()
    assert cuota_marzo.estado == Cuota.ESTADO_PAGADA
    assert cuota_abril.estado == Cuota.ESTADO_PENDIENTE


@pytest.mark.django_db
def test_cobros_gestion_rechaza_saltar_cuota_mas_vieja(client):
    staff = crear_usuario_gestion("staff_cobro_salta_cuota")
    asociado = create_asociado(
        nombre="Ana", apellido="Ferreyra", dni="49999111", tipo="asociado", fecha_alta="2026-03-10"
    )
    ciclo = CicloLectivo.objects.get_or_create(anio=2026)[0]
    marzo = PeriodoCuota.objects.create(
        mes=3,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-03-10",
    )
    abril = PeriodoCuota.objects.create(
        mes=4,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-04-10",
    )
    generar_cuotas_para_periodo(marzo)
    generar_cuotas_para_periodo(abril)
    cuota_marzo, cuota_abril = asociado.cuotas.order_by("periodo__mes")

    client.force_login(staff)
    response = client.post(
        reverse("gestion:cobros"),
        {
            "asociado_id": asociado.id,
            "cuotas_ids": [str(cuota_abril.id)],
            "fecha": "2026-04-05",
            "importe": "3000.00",
            "metodo": Pago.METODO_EFECTIVO,
            "observaciones": "",
        },
    )

    assert response.status_code == 200
    assert "deuda más antigua" in response.content.decode()
    assert Pago.objects.count() == 0
    cuota_marzo.refresh_from_db()
    cuota_abril.refresh_from_db()
    assert cuota_marzo.estado == Cuota.ESTADO_PENDIENTE
    assert cuota_abril.estado == Cuota.ESTADO_PENDIENTE


@pytest.mark.django_db
def test_deudores_gestion_lista_asociados_y_linkea_a_cobro(client):
    staff = crear_usuario_gestion("staff_deudores")
    asociado = create_asociado(
        nombre="Lucia", apellido="Ramos", dni="47777111", tipo="asociado", fecha_alta="2026-05-10"
    )
    periodo = PeriodoCuota.objects.create(
        mes=timezone.localdate().month,
        ciclo_lectivo=CicloLectivo.objects.get_or_create(anio=timezone.localdate().year)[0],
        importe="3000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
    )
    generar_cuotas_para_periodo(periodo)

    client.force_login(staff)
    response = client.get(reverse("gestion:deudores"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Ramos" in content
    assert f"?asociado={asociado.id}" in content


@pytest.mark.django_db
def test_periodos_cuota_gestion_crea_periodo(client):
    staff = crear_usuario_gestion("staff_periodo")
    client.force_login(staff)

    ciclo, _ = CicloLectivo.objects.get_or_create(anio=2026)
    response = client.post(
        reverse("gestion:periodos_cuota"),
        {
            "action": "crear_periodo",
            "mes": 6,
            "ciclo_lectivo": ciclo.id,
            "importe": "3200.00",
            "importe_recargo_mes": "500.00",
            "importe_recargo_mes_siguiente": "500.00",
            "fecha_vencimiento": "2026-06-10",
            "activo": "on",
        },
        follow=True,
    )

    assert response.status_code == 200
    assert PeriodoCuota.objects.filter(mes=6, ciclo_lectivo__anio=2026, importe="3200.00").exists()
    assert "Periodo 06/2026 creado correctamente" in response.content.decode()


@pytest.mark.django_db
def test_periodos_cuota_gestion_muestra_alerta_de_errores(client):
    staff = crear_usuario_gestion("staff_periodo_errores")
    client.force_login(staff)

    response = client.post(
        reverse("gestion:periodos_cuota"),
        {"action": "crear_periodo"},
    )

    assert response.status_code == 200
    content = response.content.decode()
    assert "Revisá los datos del período" in content
    assert "Hay campos incompletos" in content
    assert "Mes:" in content


@pytest.mark.django_db
def test_periodos_cuota_gestion_genera_cuotas_sin_duplicar(client):
    staff = crear_usuario_gestion("staff_generacion")
    create_asociado(nombre="Lara", apellido="Suarez", dni="48888111", tipo="asociado", fecha_alta="2026-05-10")
    create_asociado(nombre="Nico", apellido="Ferreyra", dni="49999111", tipo="asociado", fecha_alta="2026-05-20")
    periodo = PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=CicloLectivo.objects.get_or_create(anio=2026)[0],
        importe="3200.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-06-10",
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:periodos_cuota"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Generar cuotas" in content
    assert "0 cuotas" in content


@pytest.mark.django_db
def test_asociados_gestion_busca_y_muestra_detalle(client):
    staff = crear_usuario_gestion("staff_asoc")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10"
    )
    anio_actual = timezone.localdate().year
    ciclo_actual = CicloLectivo.objects.create(anio=anio_actual)
    ciclo_anterior = CicloLectivo.objects.create(anio=anio_actual - 1)
    periodo_actual = PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=ciclo_actual,
        importe="12000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
        activo=True,
    )
    periodo_actual_pagada = PeriodoCuota.objects.create(
        mes=5,
        ciclo_lectivo=ciclo_actual,
        importe="12000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
        activo=True,
    )
    periodo_anterior = PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=ciclo_anterior,
        importe="11000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
        activo=True,
    )
    Cuota.objects.create(asociado=asociado, periodo=periodo_actual, importe="12000.00", importe_recargo_mes="0.00", importe_recargo_mes_siguiente="0.00")
    Cuota.objects.create(
        asociado=asociado,
        periodo=periodo_actual_pagada,
        importe="12000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        importe_pagado="12000.00",
        estado=Cuota.ESTADO_PAGADA,
    )
    Cuota.objects.create(asociado=asociado, periodo=periodo_anterior, importe="11000.00", importe_recargo_mes="0.00", importe_recargo_mes_siguiente="0.00")

    client.force_login(staff)
    listado = client.get(reverse("gestion:asociados"), {"q": "Campos"})

    assert listado.status_code == 200
    content = listado.content.decode()
    assert "Campos" in content
    assert reverse("gestion:exportar_asociados") in content
    assert "?q=Campos" in content

    detalle = client.get(reverse("gestion:asociado_detalle", args=[asociado.id]))

    assert detalle.status_code == 200
    content = detalle.content.decode()
    assert "Julia" in content
    assert "40000111" in content  # usuario = DNI
    assert f"?asociado={asociado.id}" in content
    assert "Editar asociado" in content
    assert "Guardar cambios" not in content
    assert "Cuotas" in content
    assert '<th class="text-end">Importe</th>' in content
    assert '<th class="text-end">Pagado</th>' in content
    assert '<th class="text-end">Saldo</th>' in content
    assert "06/2025" not in content
    assert "05/2026" in content
    assert "Pagada" in content
    assert "Ver todas las cuotas" not in content
    assert "Ver auditoría" not in content
    assert "Crear usuario" not in content


@pytest.mark.django_db
def test_asociado_cuotas_muestra_historico_completo(client):
    staff = crear_usuario_gestion("staff_historial")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10"
    )
    anio_actual = timezone.localdate().year
    ciclo_actual = CicloLectivo.objects.create(anio=anio_actual)
    ciclo_anterior = CicloLectivo.objects.create(anio=anio_actual - 1)
    periodo_actual = PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=ciclo_actual,
        importe="12000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
        activo=True,
    )
    periodo_anterior = PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=ciclo_anterior,
        importe="11000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento=timezone.localdate(),
        activo=True,
    )
    Cuota.objects.create(asociado=asociado, periodo=periodo_actual, importe="12000.00", importe_recargo_mes="0.00", importe_recargo_mes_siguiente="0.00")
    Cuota.objects.create(asociado=asociado, periodo=periodo_anterior, importe="11000.00", importe_recargo_mes="0.00", importe_recargo_mes_siguiente="0.00")

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_cuotas", args=[asociado.id]))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Cuotas de Campos, Julia" in content
    assert "06/2025" in content
    assert "06/2026" in content
    assert "Volver al detalle" not in content


@pytest.mark.django_db
def test_asociados_gestion_busqueda_sin_filtros_lista_todos(client):
    staff = crear_usuario_gestion("staff_asoc_sin_filtros")
    create_asociado(nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10")
    create_asociado(nombre="Mario", apellido="Rivas", dni="40000112", tipo="adherente", fecha_alta="2026-05-10")

    client.force_login(staff)
    response = client.get(
        reverse("gestion:asociados"),
        {"q": "", "estado": "", "tipo": "", "curso_actual": "", "usuario": "", "deuda": ""},
    )

    assert response.status_code == 200
    content = response.content.decode()
    assert "Campos" in content
    assert "Rivas" in content
    assert "No se encontraron asociados" not in content


@pytest.mark.django_db
def test_fila_de_asociado_es_un_unico_enlace_y_conserva_la_busqueda(client):
    staff = crear_usuario_gestion("staff_fila_asociado")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000131", tipo="asociado", fecha_alta="2026-05-10"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociados"), {"q": "Campos"})

    content = response.content.decode()
    assert content.count("uni2-row-link") == 1
    assert (
        f'href="{reverse("gestion:asociado_detalle", args=[asociado.id])}'
        '?volver=/gestion/asociados/%3Fq%3DCampos"'
    ) in content
    assert "<th>Acciones</th>" not in content
    assert "Ver detalle" not in content
    assert ">Cobrar</a>" not in content


@pytest.mark.django_db
def test_detalle_conserva_retorno_filtrado_en_edicion_y_cobro(client):
    staff = crear_usuario_gestion("staff_retorno_detalle")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000132", tipo="asociado", fecha_alta="2026-05-10"
    )
    return_url = f"{reverse('gestion:asociados')}?q=Campos&estado=activo"

    client.force_login(staff)
    response = client.get(
        reverse("gestion:asociado_detalle", args=[asociado.id]),
        {"volver": return_url},
    )

    assert response.context["return_url"] == return_url
    assert parse_qs(urlsplit(response.context["edit_url"]).query)["volver"] == [return_url]
    assert parse_qs(urlsplit(response.context["cobro_url"]).query)["volver"] == [return_url]
    content = response.content.decode()
    assert "Atención al asociado" in content
    assert "Editar asociado" in content
    assert "Registrar donación" in content
    assert "Ver todas las cuotas" not in content
    assert "Ver auditoría" not in content
    assert "Crear usuario" not in content


@pytest.mark.django_db
def test_cancelar_edicion_y_cobro_vuelve_al_detalle_con_retorno(client):
    staff = crear_usuario_gestion("staff_cancelar_operacion")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000133", tipo="asociado", fecha_alta="2026-05-10"
    )
    return_url = f"{reverse('gestion:asociados')}?q=Campos"

    client.force_login(staff)
    edicion = client.get(
        reverse("gestion:asociado_editar", args=[asociado.id]),
        {"volver": return_url},
    )
    cobro = client.get(
        reverse("gestion:cobros"),
        {"asociado": asociado.id, "volver": return_url},
    )

    assert "Volver al detalle" not in edicion.content.decode()
    assert edicion.context["detail_url"] in edicion.content.decode().replace("&amp;", "&")
    assert cobro.context["detail_url"] in cobro.content.decode().replace("&amp;", "&")
    assert "Buscar asociado" not in cobro.content.decode()
    assert "Ver detalle" not in cobro.content.decode()


@pytest.mark.django_db
def test_asociado_detalle_oculta_edicion_sin_permiso(client):
    staff = crear_usuario_gestion("staff_solo_detalle", permisos=[GESTION_CONSULTAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_detalle", args=[asociado.id]))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Estado general" in content
    assert "Editar asociado" not in content
    assert "Guardar cambios" not in content


@pytest.mark.django_db
def test_asociado_detalle_muestra_a_que_corresponde_pago_reciente(client):
    staff = crear_usuario_gestion("staff_detalle_pago")
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000112", tipo="asociado", fecha_alta="2026-03-10"
    )
    ciclo = CicloLectivo.objects.get_or_create(anio=2026)[0]
    periodo = PeriodoCuota.objects.create(
        mes=3,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="0.00",
        importe_recargo_mes_siguiente="0.00",
        fecha_vencimiento="2026-03-10",
    )
    generar_cuotas_para_periodo(periodo)
    cuota = asociado.cuotas.get(periodo=periodo)
    registrar_pago(
        asociado=asociado,
        fecha=date(2026, 3, 5),
        importe=Decimal("3500.00"),
        metodo=Pago.METODO_EFECTIVO,
        cuotas_ids=[cuota.id],
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_detalle", args=[asociado.id]))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Pagos recientes" in content
    assert "Cuotas: 03/2026" in content
    assert "Donación: $ 500,00" in content


@pytest.mark.django_db
def test_asociado_editar_requiere_permiso(client):
    staff = crear_usuario_gestion("staff_solo_lectura_editar", permisos=[GESTION_CONSULTAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_editar", args=[asociado.id]))

    assert response.status_code == 403


@pytest.mark.django_db
def test_asociado_editar_muestra_formulario_separado(client):
    staff = crear_usuario_gestion("staff_edita_pantalla", permisos=[GESTION_EDITAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Julia", apellido="Campos", dni="40000111", tipo="asociado", fecha_alta="2026-05-10"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_editar", args=[asociado.id]))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Editar asociado" in content
    assert "Guardar cambios" in content
    assert reverse("gestion:asociado_detalle", args=[asociado.id]) in content


@pytest.mark.django_db
def test_asociados_gestion_oculta_acciones_sin_permiso(client):
    atencion = crear_usuario_gestion(
        "atencion_mutual",
        permisos=[
            GESTION_CONSULTAR_ASOCIADOS,
            GESTION_EDITAR_ASOCIADOS,
            GESTION_COBRAR_CUOTAS,
        ],
    )

    client.force_login(atencion)
    response = client.get(reverse("gestion:asociados"))

    assert response.status_code == 200
    content = response.content.decode()
    assert "Exportar asociados" not in content
    assert "Importar padrón inicial" not in content
    assert "Nuevo asociado" in content


@pytest.mark.django_db
def test_asociado_nuevo_crea_asociado_desde_gestion(client):
    staff = crear_usuario_gestion("staff_alta_asoc", permisos=[GESTION_CONSULTAR_ASOCIADOS, GESTION_EDITAR_ASOCIADOS])
    curso = Curso.objects.create(anio="1ro", curso="1ra", division=Curso.DIVISION_CB, turno=Curso.TURNO_TM)

    client.force_login(staff)
    response = client.post(
        reverse("gestion:asociado_nuevo"),
        {
            "nombre": "Mara",
            "apellido": "Lopez",
            "dni": "44111222",
            "email": "mara@example.com",
            "telefono": "2984000111",
            "direccion": "San Martin 100",
            "tipo": "asociado",
            "curso_actual": curso.id,
            "fecha_alta": "2026-05-20",
        },
        follow=True,
    )

    assert response.status_code == 200
    asociado = Asociado.objects.get(dni="44111222")
    assert asociado.nombre == "Mara"
    assert asociado.curso_actual == curso
    assert str(asociado.fecha_inicio_cobro) == "2026-06-01"
    assert response.redirect_chain[-1][0] == reverse("gestion:asociado_detalle", args=[asociado.id])
    assert "Asociado creado correctamente" in response.content.decode()


@pytest.mark.django_db
def test_asociado_nuevo_genera_cuotas_iniciales_y_redirige_al_detalle_aunque_pueda_cobrar(client):
    staff = crear_usuario_gestion(
        "staff_alta_cobra",
        permisos=[GESTION_CONSULTAR_ASOCIADOS, GESTION_EDITAR_ASOCIADOS, GESTION_COBRAR_CUOTAS],
    )
    curso = Curso.objects.create(anio="1ro", curso="1ra", division=Curso.DIVISION_CB, turno=Curso.TURNO_TM)
    ciclo = CicloLectivo.objects.create(anio=2026)
    PeriodoCuota.objects.create(
        mes=5,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-05-10",
    )
    PeriodoCuota.objects.create(
        mes=6,
        ciclo_lectivo=ciclo,
        importe="3000.00",
        importe_recargo_mes="500.00",
        importe_recargo_mes_siguiente="500.00",
        fecha_vencimiento="2026-06-10",
    )

    client.force_login(staff)
    response = client.post(
        reverse("gestion:asociado_nuevo"),
        {
            "nombre": "Mara",
            "apellido": "Lopez",
            "dni": "44111223",
            "email": "mara@example.com",
            "telefono": "2984000111",
            "direccion": "San Martin 100",
            "tipo": "asociado",
            "curso_actual": curso.id,
            "fecha_alta": "2026-06-05",
        },
        follow=True,
    )

    asociado = Asociado.objects.get(dni="44111223")
    assert response.status_code == 200
    assert response.redirect_chain[-1][0] == reverse("gestion:asociado_detalle", args=[asociado.id])
    assert list(asociado.cuotas.order_by("periodo__mes").values_list("periodo__mes", flat=True)) == [5, 6]
    content = response.content.decode()
    assert "Se generaron 2 cuotas iniciales" in content
    assert "Cobrar" in content


@pytest.mark.django_db
def test_asociado_nuevo_requiere_permiso_editar(client):
    staff = crear_usuario_gestion("staff_solo_consulta", permisos=[GESTION_CONSULTAR_ASOCIADOS])

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_nuevo"))

    assert response.status_code == 403


@pytest.mark.django_db
def test_asociados_gestion_filtra_por_estado_tipo_usuario_y_curso(client):
    staff = crear_usuario_gestion("staff_filtros")
    curso = Curso.objects.create(anio="2do", curso="1ra", division=Curso.DIVISION_CB, turno=Curso.TURNO_TM)
    asociado_con_usuario = create_asociado(
        nombre="Julia",
        apellido="Campos",
        dni="40000111",
        tipo="asociado",
        fecha_alta="2026-05-10",
        curso_actual=curso,
    )
    user_model = get_user_model()
    user = user_model.objects.create_user(username="julia", password="secreto123")
    asociado_con_usuario.usuario = user
    asociado_con_usuario.save(update_fields=["usuario"])
    create_asociado(
        nombre="Mora",
        apellido="Rivas",
        dni="40000222",
        tipo="adherente",
        fecha_alta="2026-05-11",
    )

    client.force_login(staff)
    response = client.get(
        reverse("gestion:asociados"),
        {
            "tipo": "asociado",
            "estado": "activo",
            "curso_actual": curso.id,
            "usuario": "con",
        },
    )

    assert response.status_code == 200
    content = response.content.decode()
    assert "Campos" in content
    assert "Rivas" not in content
    assert "Con usuario" in content
    assert "Asociado" in content
    assert "Todos los cursos" in content


@pytest.mark.django_db
def test_exportar_asociados_descarga_formato_uni2_filtrado(client):
    staff = crear_usuario_gestion("staff_exporta")
    curso = Curso.objects.create(anio="1ro", curso="2da", division=Curso.DIVISION_CB, turno=Curso.TURNO_TM)
    create_asociado(
        nombre="Julia",
        apellido="Campos",
        dni="40000111",
        tipo="asociado",
        fecha_alta="2026-05-10",
        curso_actual=curso,
        email="julia@example.com",
        telefono="2984 123456",
        direccion="Calle 1",
    )
    create_asociado(
        nombre="Mora",
        apellido="Rivas",
        dni="40000222",
        tipo="adherente",
        fecha_alta="2026-05-11",
    )

    client.force_login(staff)
    response = client.get(
        reverse("gestion:exportar_asociados"),
        {
            "tipo": "asociado",
            "estado": "activo",
            "curso_actual": curso.id,
            "usuario": "con",
        },
    )

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "asociados_formato_uni2.xlsx" in response["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["ASOCIADOS"]
    assert [worksheet.cell(1, column).value for column in range(1, 18)] == [
        "numero_asociado",
        "apellido",
        "nombre",
        "dni",
        "tipo",
        "email",
        "telefono",
        "direccion",
        "curso_anio",
        "curso_division",
        "division",
        "turno",
        "curso_nombre",
        "fecha_alta",
        "fecha_inicio_cobro",
        "estado",
        "motivo_baja",
    ]
    assert worksheet.max_row == 2
    row_values = [worksheet.cell(2, column).value for column in range(1, 18)]
    assert row_values[1:17] == [
        "Campos",
        "Julia",
        "40000111",
        "asociado",
        "julia@example.com",
        "2984 123456",
        "Calle 1",
        "1ro",
        "2da",
        "CB",
        "TM",
        "1ro 2da CB TM",
        "2026-05-10",
        "2026-05-01",
        "activo",
        None,
    ]


@pytest.mark.django_db
def test_asociado_detalle_permite_editar_fecha_inicio_cobro(client):
    staff = crear_usuario_gestion("staff_edita_asoc")
    asociado = create_asociado(
        nombre="Milena", apellido="Armada", dni="30000111", tipo="asociado", fecha_alta="2026-05-22"
    )

    return_url = f"{reverse('gestion:asociados')}?q=Armada"
    client.force_login(staff)
    response = client.post(
        reverse("gestion:asociado_editar", args=[asociado.id]),
        {
            "nombre": "Milena",
            "apellido": "Armada",
            "dni": "30000111",
            "email": "milena@example.com",
            "telefono": "123456",
            "tipo": "asociado",
            "curso_actual": "",
            "estado": "activo",
            "fecha_alta": "2026-05-22",
            "fecha_inicio_cobro": "2026-05-01",
            "fecha_baja": "",
            "motivo_baja": "",
            "volver": return_url,
        },
        follow=True,
    )

    assert response.status_code == 200
    asociado.refresh_from_db()
    assert str(asociado.fecha_inicio_cobro) == "2026-05-01"
    assert asociado.email == "milena@example.com"
    assert "Asociado actualizado correctamente" in response.content.decode()
    redirect_url = response.redirect_chain[-1][0]
    assert urlsplit(redirect_url).path == reverse("gestion:asociado_detalle", args=[asociado.id])
    assert parse_qs(urlsplit(redirect_url).query)["volver"] == [return_url]


@pytest.mark.django_db
def test_edicion_cotidiana_no_expone_ni_modifica_la_baja(client):
    staff = crear_usuario_gestion("staff_sin_baja")
    asociado = create_asociado(
        nombre="Mila", apellido="Ríos", dni="30000112", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_editar", args=[asociado.id]))

    assert response.status_code == 200
    assert "fecha_baja" not in response.context["form"].fields
    assert "motivo_baja" not in response.context["form"].fields
    assert "estado" not in response.context["form"].fields

    client.post(
        reverse("gestion:asociado_editar", args=[asociado.id]),
        {
            "nombre": asociado.nombre,
            "apellido": asociado.apellido,
            "dni": asociado.dni,
            "email": asociado.email,
            "telefono": asociado.telefono,
            "direccion": asociado.direccion,
            "tipo": asociado.tipo,
            "curso_actual": "",
            "estado": Asociado.ESTADO_INACTIVO,
            "fecha_alta": asociado.fecha_alta,
            "fecha_inicio_cobro": asociado.fecha_inicio_cobro,
            "fecha_baja": "2026-08-09",
            "motivo_baja": "No debe aplicarse",
        },
    )

    asociado.refresh_from_db()
    assert asociado.estado == Asociado.ESTADO_ACTIVO
    assert asociado.fecha_baja is None
    assert asociado.motivo_baja == ""


@pytest.mark.django_db
def test_crear_usuario_asociado_requiere_permiso_editar(client):
    user = get_user_model().objects.create_user(username="sin_permiso", password="secreto123")
    asociado = create_asociado(
        nombre="Juan", apellido="Perez", dni="30000222", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(user)
    response = client.post(reverse("gestion:crear_usuario_asociado", args=[asociado.id]), {"password": "test123"})

    assert response.status_code == 403


@pytest.mark.django_db
def test_crear_usuario_asociado_crea_y_vincula(client):
    staff = crear_usuario_gestion("staff_crea_user", permisos=[GESTION_EDITAR_ASOCIADOS, GESTION_CONSULTAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Juan", apellido="Perez", dni="30000222", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(staff)
    response = client.post(reverse("gestion:crear_usuario_asociado", args=[asociado.id]), follow=True)

    assert response.status_code == 200
    asociado.refresh_from_db()
    assert asociado.usuario is not None
    assert asociado.usuario.username == "30000222"
    assert "ya tiene un usuario" in response.content.decode()


@pytest.mark.django_db
def test_crear_usuario_asociado_ya_tiene_usuario_muestra_error(client):
    staff = crear_usuario_gestion("staff_crea_user2", permisos=[GESTION_EDITAR_ASOCIADOS, GESTION_CONSULTAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Juan", apellido="Perez", dni="30000222", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(staff)
    response = client.post(reverse("gestion:crear_usuario_asociado", args=[asociado.id]), follow=True)

    assert response.status_code == 200
    assert "ya tiene un usuario" in response.content.decode()


@pytest.mark.django_db
def test_crear_usuario_asociado_get_redirige_a_detalle(client):
    staff = crear_usuario_gestion("staff_crea_user3", permisos=[GESTION_EDITAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Juan", apellido="Perez", dni="30000222", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:crear_usuario_asociado", args=[asociado.id]))

    assert response.status_code == 302
    assert reverse("gestion:asociado_detalle", args=[asociado.id]) in response.url


@pytest.mark.django_db
def test_asociado_detalle_oculta_boton_crear_usuario_cuando_ya_tiene_usuario(client):
    staff = crear_usuario_gestion("staff_no_ve_boton", permisos=[GESTION_EDITAR_ASOCIADOS, GESTION_CONSULTAR_ASOCIADOS])
    asociado = create_asociado(
        nombre="Juan", apellido="Perez", dni="30000222", tipo="asociado", fecha_alta="2026-05-22"
    )

    client.force_login(staff)
    response = client.get(reverse("gestion:asociado_detalle", args=[asociado.id]))

    assert response.status_code == 200
    assert "Crear usuario" not in response.content.decode()
