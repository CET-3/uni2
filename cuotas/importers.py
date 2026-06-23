from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from django.db import transaction

from asociados.models import Asociado, CicloLectivo

from .models import Cuota, Pago, PagoCuota, PeriodoCuota


CUOTAS_HISTORICAS_SESSION_KEY = "cuotas_historicas_preview"
CUOTAS_HISTORICAS_SHEET = "COBRO CUOTAS SOCIALES"
CUOTAS_HISTORICAS_HEADER_ROW = 7
CUOTAS_HISTORICAS_FIRST_DATA_ROW = 8
CUOTA_HISTORICA_IMPORTE_INICIAL = Decimal("500.00")
CUOTA_HISTORICA_RECARGO_INICIAL = Decimal("100.00")
CUOTA_HISTORICA_IMPORTE_DESDE_MAYO = Decimal("600.00")
CUOTA_HISTORICA_RECARGO_DESDE_MAYO = Decimal("200.00")
CUOTA_HISTORICA_ANIO = 2026
CUOTAS_HISTORICAS_LOG_EVERY = 25
logger = logging.getLogger(__name__)
CUOTAS_HISTORICAS_REVISAR_HEADERS = [
    "Fila original",
    "Número de asociado",
    "Apellido/Nombre",
    "curso/división",
]
CUOTA_HISTORICA_MESES = [
    {"mes": 3, "nombre": "Mar", "pago_col": 4, "forma_col": 5},
    {"mes": 4, "nombre": "Abr", "pago_col": 6, "forma_col": 7},
    {"mes": 5, "nombre": "May", "pago_col": 8, "forma_col": 9},
    {"mes": 6, "nombre": "Jun", "pago_col": 10, "forma_col": 11},
    {"mes": 7, "nombre": "Jul", "pago_col": 12, "forma_col": 13},
    {"mes": 8, "nombre": "Agos", "pago_col": 14, "forma_col": 15},
    {"mes": 9, "nombre": "Sept", "pago_col": 16, "forma_col": 17},
    {"mes": 10, "nombre": "Oct", "pago_col": 18, "forma_col": 19},
    {"mes": 11, "nombre": "Nov", "pago_col": 20, "forma_col": 21},
    {"mes": 12, "nombre": "DIC", "pago_col": 22, "forma_col": 23},
]


@dataclass
class CuotasHistoricasPreview:
    importables: list[dict[str, Any]] = field(default_factory=list)
    revisar: list[dict[str, Any]] = field(default_factory=list)
    omitidas: int = 0

    @property
    def summary(self):
        pagos_a_crear = sum(1 for item in self.importables if item["pagada"])
        cuotas_impagas = sum(1 for item in self.importables if not item["pagada"])
        return {
            "cuotas_a_crear": len(self.importables),
            "pagos_a_crear": pagos_a_crear,
            "cuotas_impagas": cuotas_impagas,
            "revisar": len(self.revisar),
            "omitidas": self.omitidas,
        }

    def as_session_data(self):
        return {
            "summary": self.summary,
            "importables": self.importables,
            "revisar": self.revisar,
            "omitidas": self.omitidas,
        }

    @classmethod
    def from_session_data(cls, data):
        return cls(
            importables=data.get("importables", []),
            revisar=data.get("revisar", []),
            omitidas=data.get("omitidas", data.get("summary", {}).get("omitidas", 0)),
        )


@dataclass
class CuotasHistoricasImportResult:
    periodos_creados: int = 0
    cuotas_creadas: int = 0
    pagos_creados: int = 0
    omitidas: int = 0
    errores: list[str] = field(default_factory=list)


def _normalize_name(value):
    text = unicodedata.normalize("NFKD", str(value).lower().strip())
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text)


def _parse_curso_parts(curso_raw):
    """Parse '1°1°' into (anio: '1ro', curso: '1ra') or (None, None)."""
    if not curso_raw:
        return None, None
    text = curso_raw.lower().strip().replace("º", "°")
    text = text.replace(".", " ").replace(",", " ")
    text = re.sub(r"\s+", " ", text).strip()
    nums = re.findall(r"\d+", text)
    if len(nums) < 2:
        return None, None
    anio_map = {"1": "1ro", "2": "2do", "3": "3ro", "4": "4to"}
    div_map = {"1": "1ra", "2": "2da", "3": "3ra", "4": "4ta"}
    return anio_map.get(nums[0]), div_map.get(nums[1])


def _build_asociados_lookup():
    """Return (by_name, by_apellido_curso) lookup dicts.

    by_apellido_curso keys are (apellido_normalized, anio, curso_division)
    ignoring ciclo (CB/CS) and turno, so the cuotas import can find
    asociados even when the planilla omits or mislabels the cycle.
    Values are lists to detect ambiguity (same apellido in same
    anio+curso).
    """
    by_name = {}
    by_apellido_curso = {}
    for a in Asociado.objects.select_related("curso_actual").all():
        key = _normalize_name(f"{a.apellido} {a.nombre}")
        by_name[key] = a
        if a.curso_actual:
            apellido_key = _normalize_name(a.apellido)
            curso_key = (apellido_key, a.curso_actual.anio, a.curso_actual.curso)
            by_apellido_curso.setdefault(curso_key, []).append(a)
    return by_name, by_apellido_curso


def _find_asociado(nombre_original, curso_original, by_name, by_apellido_curso):
    """Find asociado by name, falling back to course+apellido if exact match fails.

    Tries these strategies in order:
      1. Normalized exact match (e.g. "Garcia Juan" → "garcia juan").
      2. Reversed order (e.g. "Juan Garcia" → "garcia juan").
      3. First word as apellido + curso.
      4. Last  word as apellido + curso.

    The curso fallback ignores ciclo (CB/CS) and turno.
    """
    name_key = _normalize_name(nombre_original)
    asociado = by_name.get(name_key)
    if asociado:
        return asociado

    parts = name_key.split()
    if len(parts) >= 2:
        reversed_key = " ".join([parts[-1]] + parts[:-1])
        asociado = by_name.get(reversed_key)
        if asociado:
            return asociado

    if not nombre_original or not curso_original:
        return None

    anio, curso_num = _parse_curso_parts(curso_original)
    if not anio or not curso_num:
        return None

    partes = nombre_original.split()
    candidates = set()
    if partes:
        candidates.add(partes[0])
        candidates.add(partes[-1])
    for raw in candidates:
        apellido_key = _normalize_name(raw)
        matches = by_apellido_curso.get((apellido_key, anio, curso_num), [])
        if len(matches) == 1:
            return matches[0]

    return None


def _clean_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_numero_asociado(value):
    text = _clean_value(value)
    if not text:
        return None
    if not text.isdigit():
        return None
    return int(text)


def _normalize_bool(value):
    if value is True or value is False:
        return value, ""
    if value is None or _clean_value(value) == "":
        return False, ""
    text = _clean_value(value).lower()
    if text in {"si", "sí", "s", "true", "1", "pagado"}:
        return True, ""
    if text in {"no", "false", "0"}:
        return False, ""
    return False, f'Valor de pago dudoso: "{_clean_value(value)}"'


def _normalize_metodo(value):
    text = _clean_value(value).lower()
    if not text:
        return "", ""
    if text in {"efectivo", "efe"}:
        return Pago.METODO_EFECTIVO, ""
    if text in {"mp", "mercado pago", "billetera", "billetera virtual"}:
        return Pago.METODO_BILLETERA, ""
    return "", f'Forma de pago dudosa: "{_clean_value(value)}"'


def _periodo_fecha(mes):
    return date(CUOTA_HISTORICA_ANIO, mes, 10)


def _valores_cuota_historica(mes):
    if mes >= 5:
        return CUOTA_HISTORICA_IMPORTE_DESDE_MAYO, CUOTA_HISTORICA_RECARGO_DESDE_MAYO
    return CUOTA_HISTORICA_IMPORTE_INICIAL, CUOTA_HISTORICA_RECARGO_INICIAL


def _estado_impaga(mes, fecha_operacion):
    if fecha_operacion > _periodo_fecha(mes):
        return Cuota.ESTADO_VENCIDA
    return Cuota.ESTADO_PENDIENTE


def analyze_cuotas_historicas_xlsx(file_obj, fecha_operacion: date) -> CuotasHistoricasPreview:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Para importar planillas .xlsx hace falta instalar openpyxl.") from exc

    workbook = load_workbook(file_obj, data_only=True)
    if CUOTAS_HISTORICAS_SHEET not in workbook.sheetnames:
        raise ValueError(f'El archivo debe tener una hoja llamada "{CUOTAS_HISTORICAS_SHEET}".')

    worksheet = workbook[CUOTAS_HISTORICAS_SHEET]
    meses = [item for item in CUOTA_HISTORICA_MESES if item["mes"] <= fecha_operacion.month]
    preview = CuotasHistoricasPreview()
    by_name, by_apellido_curso = _build_asociados_lookup()

    for row_number in range(CUOTAS_HISTORICAS_FIRST_DATA_ROW, worksheet.max_row + 1):
        numero = _clean_numero_asociado(worksheet.cell(row_number, 1).value)
        nombre_original = _clean_value(worksheet.cell(row_number, 2).value)
        curso_original = _clean_value(worksheet.cell(row_number, 3).value)
        if not nombre_original:
            preview.omitidas += 1
            continue

        asociado = _find_asociado(nombre_original, curso_original, by_name, by_apellido_curso)
        for mes_data in meses:
            raw_pagada = worksheet.cell(row_number, mes_data["pago_col"]).value
            raw_forma = worksheet.cell(row_number, mes_data["forma_col"]).value
            pagada, pago_note = _normalize_bool(raw_pagada)
            metodo, metodo_note = _normalize_metodo(raw_forma)
            if not pagada and metodo:
                pagada = True
            observaciones = [item for item in [pago_note, metodo_note] if item]
            importe, recargo = _valores_cuota_historica(mes_data["mes"])

            item = {
                "fila_origen": row_number,
                "numero_asociado": numero,
                "apellido_nombre_original": nombre_original,
                "curso_original": curso_original,
                "mes": mes_data["mes"],
                "mes_nombre": mes_data["nombre"],
                "pagada": pagada,
                "forma_pago_original": _clean_value(raw_forma),
                "metodo": metodo,
                "asociado_id": asociado.id if asociado else None,
                "asociado_nombre": str(asociado) if asociado else "",
                "importe": str(importe),
                "importe_recargo_mes": str(recargo),
                "importe_recargo_mes_siguiente": str(recargo),
                "estado": Cuota.ESTADO_PAGADA if pagada else _estado_impaga(mes_data["mes"], fecha_operacion),
            }

            if not asociado:
                observaciones.append("No se encontró asociado por nombre")
            if pagada and not metodo:
                metodo = Pago.METODO_EFECTIVO
            if asociado and PeriodoCuota.objects.filter(
                mes=mes_data["mes"], ciclo_lectivo__anio=CUOTA_HISTORICA_ANIO, cuotas__asociado=asociado
            ).exists():
                observaciones.append("La cuota ya existe para este asociado y período")

            if observaciones:
                item["observaciones"] = "; ".join(dict.fromkeys(observaciones))
                preview.revisar.append(item)
            else:
                preview.importables.append(item)

    return preview


def build_revisar_cuotas_historicas_xlsx(preview: CuotasHistoricasPreview) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Para descargar planillas .xlsx hace falta instalar openpyxl.") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = CUOTAS_HISTORICAS_SHEET
    months_by_number = {item["mes"]: item["nombre"] for item in CUOTA_HISTORICA_MESES}
    months_with_errors = sorted({int(row["mes"]) for row in preview.revisar})
    headers = CUOTAS_HISTORICAS_REVISAR_HEADERS[:]
    for mes in months_with_errors:
        month_name = months_by_number[mes]
        headers.extend([f"{month_name} pago", f"{month_name} forma", f"{month_name} motivo"])
    worksheet.append(headers)

    grouped_rows = {}
    for row in preview.revisar:
        key = (row.get("fila_origen"), row.get("numero_asociado"), row.get("apellido_nombre_original"))
        if key not in grouped_rows:
            grouped_rows[key] = {
                "fila_origen": row.get("fila_origen", ""),
                "numero_asociado": row.get("numero_asociado", ""),
                "apellido_nombre_original": row.get("apellido_nombre_original", ""),
                "curso_original": row.get("curso_original", ""),
                "meses": {},
            }
        grouped_rows[key]["meses"][int(row["mes"])] = {
            "pago": "Sí" if row.get("pagada") else "No",
            "forma": row.get("forma_pago_original", ""),
            "motivo": row.get("observaciones", ""),
        }

    for row in sorted(grouped_rows.values(), key=lambda item: item["fila_origen"]):
        output_row = [
            row["fila_origen"],
            row["numero_asociado"],
            row["apellido_nombre_original"],
            row["curso_original"],
        ]
        for mes in months_with_errors:
            month_data = row["meses"].get(mes, {})
            output_row.extend([month_data.get("pago", ""), month_data.get("forma", ""), month_data.get("motivo", "")])
        worksheet.append(output_row)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@transaction.atomic
def import_cuotas_historicas_preview(preview: CuotasHistoricasPreview, registrado_por=None):
    result = CuotasHistoricasImportResult()
    ciclo, _ = CicloLectivo.objects.get_or_create(anio=CUOTA_HISTORICA_ANIO)
    periodos = {}
    total = len(preview.importables)
    logger.info("Importacion de cuotas historicas iniciada: %s cuotas importables.", total)

    for index, item in enumerate(preview.importables, start=1):
        try:
            mes = int(item["mes"])
            importe, recargo = _valores_cuota_historica(mes)
            periodo_key = (CUOTA_HISTORICA_ANIO, mes)
            if periodo_key not in periodos:
                periodo, created = PeriodoCuota.objects.get_or_create(
                    mes=mes,
                    ciclo_lectivo=ciclo,
                    defaults={
                        "importe": importe,
                        "importe_recargo_mes": recargo,
                        "importe_recargo_mes_siguiente": recargo,
                        "fecha_vencimiento": _periodo_fecha(mes),
                        "activo": True,
                    },
                )
                periodos[periodo_key] = periodo
                result.periodos_creados += int(created)
            periodo = periodos[periodo_key]

            asociado = Asociado.objects.get(id=item["asociado_id"])
            cuota, created = Cuota.objects.get_or_create(
                asociado=asociado,
                periodo=periodo,
                defaults={
                    "importe": importe,
                    "importe_recargo_mes": recargo,
                    "importe_recargo_mes_siguiente": recargo,
                    "importe_pagado": importe if item["pagada"] else Decimal("0.00"),
                    "estado": item["estado"],
                },
            )
            if not created:
                result.omitidas += 1
                continue
            result.cuotas_creadas += 1

            if item["pagada"]:
                pago = Pago.objects.create(
                    asociado=asociado,
                    fecha=_periodo_fecha(mes),
                    importe=importe,
                    metodo=item["metodo"],
                    registrado_por=registrado_por,
                    observaciones=(
                        f"Importado desde planilla historica de cuotas. "
                        f"Fila original {item['fila_origen']}, mes {item['mes_nombre']}."
                    ),
                )
                PagoCuota.objects.create(pago=pago, cuota=cuota, importe=importe)
                result.pagos_creados += 1
        except Exception as exc:  # noqa: BLE001
            result.errores.append(f"Fila {item.get('fila_origen')} mes {item.get('mes_nombre')}: {exc}")

        if index == total or index % CUOTAS_HISTORICAS_LOG_EVERY == 0:
            logger.info("Importacion de cuotas historicas: %s/%s cuotas procesadas.", index, total)

    result.omitidas += preview.omitidas
    logger.info(
        "Importacion de cuotas historicas finalizada: %s periodos creados, %s cuotas creadas, "
        "%s pagos creados, %s omitidas, %s errores.",
        result.periodos_creados,
        result.cuotas_creadas,
        result.pagos_creados,
        result.omitidas,
        len(result.errores),
    )
    return result
