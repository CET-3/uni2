from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import BinaryIO

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.validators import URLValidator
from django.db import transaction
from django.utils.text import slugify
from PIL import Image as PillowImage

from .models import ActividadComercial, Comercio


COMERCIOS_HEADERS = {
    "COMERCIO": "nombre",
    "ACTIVIDAD COMERCIAL": "actividad_comercial",
    "DESCUENTO": "beneficio_texto",
    "INSTAGRAM": "url_presencia_web",
    "UBICACION": "direccion",
    "DESCRIPCION": "descripcion",
    "LOGO": "logo",
}


@dataclass(frozen=True)
class ComercioImportImage:
    content: bytes
    extension: str
    digest: str


@dataclass(frozen=True)
class ComercioImportRow:
    fila_origen: int
    nombre: str
    actividad_comercial: str
    beneficio_texto: str
    url_presencia_web: str
    direccion: str
    descripcion: str
    foto: ComercioImportImage | None = None


@dataclass
class ComerciosImportPreview:
    hoja: str
    filas: list[ComercioImportRow] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)

    @property
    def cantidad_actividades(self):
        activities = {
            _identity(row.actividad_comercial)
            for row in self.filas
            if row.actividad_comercial
        }
        return len(activities)

    @property
    def cantidad_fotos(self):
        return sum(1 for row in self.filas if row.foto is not None)

    @property
    def es_valida(self):
        return not self.errores


@dataclass
class ComerciosImportResult:
    comercios_creados: int = 0
    comercios_actualizados: int = 0
    actividades_creadas: int = 0
    fotos_asignadas: int = 0


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().upper())
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text)


def _identity(value):
    text = " ".join(str(value or "").strip().casefold().split())
    normalized = unicodedata.normalize("NFKD", text)
    return normalized.encode("ascii", "ignore").decode("ascii")


def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).strip())


def _clean_beneficio(value):
    if isinstance(value, bool):
        return _clean_text(value)
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if Decimal("0") <= number <= Decimal("1"):
            percentage = (number * 100).normalize()
            return f"{format(percentage, 'f')}%"
    return _clean_text(value)


def _header_columns(worksheet):
    columns = {}
    repeated = []
    for column in range(1, worksheet.max_column + 1):
        normalized = _normalize_header(worksheet.cell(1, column).value)
        if normalized in COMERCIOS_HEADERS:
            if normalized in columns:
                repeated.append(normalized)
            columns[normalized] = column

    if repeated:
        names = ", ".join(sorted(set(repeated)))
        raise ValueError(f"La fila de encabezados tiene columnas repetidas: {names}.")

    missing = sorted(set(COMERCIOS_HEADERS) - set(columns))
    if missing:
        names = ", ".join(missing)
        raise ValueError(f"Faltan columnas obligatorias en la fila de encabezados: {names}.")

    return {COMERCIOS_HEADERS[header]: column for header, column in columns.items()}


def _extract_image(image):
    try:
        content = image._data()
        with PillowImage.open(BytesIO(content)) as opened:
            image_format = (opened.format or "").lower()
            opened.verify()
    except Exception as exc:
        raise ValueError("la imagen embebida no tiene un formato válido") from exc

    extension = {"jpeg": "jpg", "jpg": "jpg", "png": "png", "webp": "webp"}.get(image_format)
    if extension is None:
        raise ValueError(f'el formato de imagen "{image_format or "desconocido"}" no está soportado')

    return ComercioImportImage(
        content=content,
        extension=extension,
        digest=hashlib.sha256(content).hexdigest(),
    )


def _extract_images_by_row(worksheet, logo_column):
    images_by_row = {}
    errors = []
    warnings = []

    for image in worksheet._images:
        marker = getattr(image.anchor, "_from", None)
        if marker is None:
            warnings.append("Se ignoró una imagen porque no se pudo determinar su celda de origen.")
            continue

        row_number = marker.row + 1
        column_number = marker.col + 1
        if column_number != logo_column:
            warnings.append(
                f"Se ignoró una imagen de la fila {row_number} porque no está anclada en la columna LOGO."
            )
            continue
        if row_number in images_by_row:
            errors.append(f"Fila {row_number}: hay más de una imagen en la columna LOGO.")
            continue

        try:
            images_by_row[row_number] = _extract_image(image)
        except ValueError as exc:
            errors.append(f"Fila {row_number}: {exc}.")

    return images_by_row, errors, warnings


def _validate_max_length(*, row_number, label, value, model_field, errors):
    if value and model_field.max_length and len(value) > model_field.max_length:
        errors.append(
            f"Fila {row_number}: {label} supera el máximo de {model_field.max_length} caracteres."
        )


def analyze_comercios_xlsx(file_obj: BinaryIO) -> ComerciosImportPreview:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Para importar comercios hace falta instalar openpyxl.") from exc

    try:
        workbook = load_workbook(file_obj, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo leer el archivo como una planilla .xlsx válida.") from exc

    worksheet = workbook.active
    columns = _header_columns(worksheet)
    images_by_row, image_errors, image_warnings = _extract_images_by_row(
        worksheet,
        columns["logo"],
    )
    preview = ComerciosImportPreview(
        hoja=worksheet.title,
        errores=image_errors,
        advertencias=image_warnings,
    )

    last_row = max([worksheet.max_row, *images_by_row.keys()])
    seen_names = {}
    url_validator = URLValidator(schemes=["http", "https"])

    for row_number in range(2, last_row + 1):
        raw_values = {
            field_name: worksheet.cell(row_number, column).value
            for field_name, column in columns.items()
            if field_name != "logo"
        }
        has_values = any(value not in (None, "") for value in raw_values.values())
        if not has_values and row_number not in images_by_row:
            continue

        row_errors = []
        nombre = _clean_text(raw_values["nombre"])
        actividad = _clean_text(raw_values["actividad_comercial"])
        beneficio = _clean_beneficio(raw_values["beneficio_texto"])
        url = _clean_text(raw_values["url_presencia_web"])
        direccion = _clean_text(raw_values["direccion"])
        descripcion = _clean_text(raw_values["descripcion"])

        required_values = (
            ("COMERCIO", nombre),
            ("ACTIVIDAD COMERCIAL", actividad),
            ("DESCUENTO", beneficio),
            ("DESCRIPCIÓN", descripcion),
        )
        for label, value in required_values:
            if not value:
                row_errors.append(f"Fila {row_number}: falta {label}.")

        if url:
            try:
                url_validator(url)
            except ValidationError:
                row_errors.append(
                    f"Fila {row_number}: INSTAGRAM debe contener una URL completa http:// o https://."
                )

        _validate_max_length(
            row_number=row_number,
            label="COMERCIO",
            value=nombre,
            model_field=Comercio._meta.get_field("nombre"),
            errors=row_errors,
        )
        _validate_max_length(
            row_number=row_number,
            label="ACTIVIDAD COMERCIAL",
            value=actividad,
            model_field=ActividadComercial._meta.get_field("nombre"),
            errors=row_errors,
        )
        _validate_max_length(
            row_number=row_number,
            label="INSTAGRAM",
            value=url,
            model_field=Comercio._meta.get_field("url_presencia_web"),
            errors=row_errors,
        )
        _validate_max_length(
            row_number=row_number,
            label="UBICACIÓN",
            value=direccion,
            model_field=Comercio._meta.get_field("direccion"),
            errors=row_errors,
        )

        name_key = _identity(nombre)
        if name_key:
            if name_key in seen_names:
                row_errors.append(
                    f'Fila {row_number}: el comercio "{nombre}" está repetido; '
                    f"también aparece en la fila {seen_names[name_key]}."
                )
            else:
                seen_names[name_key] = row_number

        preview.filas.append(
            ComercioImportRow(
                fila_origen=row_number,
                nombre=nombre,
                actividad_comercial=actividad,
                beneficio_texto=beneficio,
                url_presencia_web=url,
                direccion=direccion,
                descripcion=descripcion,
                foto=images_by_row.get(row_number),
            )
        )
        preview.errores.extend(row_errors)

    if not preview.filas:
        preview.errores.append("La planilla no contiene filas de comercios.")

    workbook.close()
    return preview


def _objects_by_identity(objects, *, label):
    indexed = {}
    for obj in objects:
        key = _identity(obj.nombre)
        if key in indexed:
            raise ValueError(
                f'La base de datos tiene más de {label} con el nombre "{obj.nombre}". '
                "Hay que corregir esa duplicación antes de importar."
            )
        indexed[key] = obj
    return indexed


def _assign_photo(comercio, image):
    image_name = slugify(comercio.nombre)[:80] or "comercio"
    filename = f"importacion/{image_name}-{image.digest[:12]}.{image.extension}"
    expected_name = comercio.foto.field.generate_filename(comercio, filename)

    if comercio.foto.name == expected_name:
        return False

    storage = comercio.foto.storage
    if storage.exists(expected_name):
        comercio.foto.name = expected_name
        comercio.save(update_fields=["foto"])
    else:
        comercio.foto.save(filename, ContentFile(image.content), save=True)
    return True


@transaction.atomic
def import_comercios_preview(preview: ComerciosImportPreview) -> ComerciosImportResult:
    if not preview.es_valida:
        raise ValueError("No se puede importar una planilla que tiene errores.")

    result = ComerciosImportResult()
    activities = _objects_by_identity(
        ActividadComercial.objects.all(),
        label="una actividad comercial",
    )
    commerces = _objects_by_identity(Comercio.objects.all(), label="un comercio")

    for row in preview.filas:
        activity_key = _identity(row.actividad_comercial)
        actividad = activities.get(activity_key)
        if actividad is None:
            actividad = ActividadComercial(nombre=row.actividad_comercial)
            actividad.full_clean()
            actividad.save()
            activities[activity_key] = actividad
            result.actividades_creadas += 1

        commerce_key = _identity(row.nombre)
        comercio = commerces.get(commerce_key)
        created = comercio is None
        if created:
            comercio = Comercio(nombre=row.nombre)

        comercio.actividad_comercial = actividad
        comercio.nombre = row.nombre
        comercio.descripcion = row.descripcion
        comercio.beneficio_texto = row.beneficio_texto
        comercio.estado = Comercio.ESTADO_FIRMADO
        comercio.orden = 0
        comercio.direccion = row.direccion
        comercio.url_presencia_web = row.url_presencia_web or None
        comercio.full_clean()
        comercio.save()

        if created:
            commerces[commerce_key] = comercio
            result.comercios_creados += 1
        else:
            result.comercios_actualizados += 1

        if row.foto is not None and _assign_photo(comercio, row.foto):
            result.fotos_asignadas += 1

    return result
